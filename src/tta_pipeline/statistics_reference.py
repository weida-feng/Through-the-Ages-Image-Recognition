"""Compare Python statistics with preserved Excel formula text and saved values."""
import hashlib
import math
from pathlib import Path
import re
from typing import Any, Dict, Tuple

import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd

from .catalog import CATALOGS, EMPTY_CARD
from .statistics import wilson_interval


def compare_statistics_reference(
    data_root: Path, tables: Dict[str, pd.DataFrame]
) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    path = Path(data_root) / "source_workbooks/5239.xlsx"
    formulas_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = []
    seen = set()
    metrics = tables["card_statistics"]
    metrics = metrics.loc[metrics.cohort.eq("historical_5239")].set_index(["kind", "card_name"])
    coverage = tables["leader_coverage"]
    coverage = coverage.loc[coverage.cohort.eq("historical_5239")].set_index("age_index")
    if set(metrics.cohort_games) != {5239}:
        raise ValueError("Historical statistics comparison requires the complete 5239-game cohort")
    try:
        formulas, cached = formulas_book["統計"], cached_book["統計"]
        for kind, row_numbers in (("leader", range(3, 16)), ("wonder", range(18, 26))):
            for age_index, base in enumerate((2, 7, 12, 17)):
                for row_number in row_numbers:
                    name = formulas.cell(row_number, base).value
                    key = (kind, age_index, name)
                    if key in seen:
                        raise ValueError("Duplicate card/age in historical statistics")
                    seen.add(key)
                    if name == EMPTY_CARD and kind == "leader":
                        absent = coverage.loc[age_index]
                        n, wins = int(absent.absent_slots), int(absent.absent_wins)
                    else:
                        if name not in CATALOGS[kind] or CATALOGS[kind][name].age != age_index:
                            raise ValueError("Unexpected historical card label {!r}".format(name))
                        card = metrics.loc[(kind, name)]
                        n, wins = int(card.usage_count), int(card.wins)
                    label = "{}{}".format(get_column_letter(base), row_number)
                    if kind == "leader":
                        win_column = get_column_letter(6 + age_index)
                        loss_column = get_column_letter(10 + age_index)
                        winning = "COUNTIF(精簡信息!{0}:{0},{1})".format(win_column, label)
                        losing = "COUNTIF(精簡信息!{0}:{0},{1})".format(loss_column, label)
                    else:
                        winning = "COUNTIF(精簡信息!$N:$U,{})".format(label)
                        losing = "COUNTIF(精簡信息!$V:$AC,{})".format(label)
                    expected_selection = "=({}+{})/5239".format(winning, losing)
                    expected_win_rate = "={}/({}+{})".format(winning, winning, losing)
                    if formulas.cell(row_number, base + 1).value != expected_selection or \
                            formulas.cell(row_number, base + 2).value != expected_win_rate:
                        raise ValueError("Historical COUNTIF formula layout differs at {}".format(label))
                    score_formula = formulas.cell(row_number, base + 3).value
                    scale_match = re.search(r"\)/(\d+(?:\.\d+)?)\*50$", str(score_formula))
                    if not scale_match or "1.96*1.96" not in score_formula:
                        raise ValueError("Historical score formula differs at {}".format(label))
                    scale = float(scale_match.group(1))
                    if scale <= 0:
                        raise ValueError("Invalid historical score scale")
                    low, _ = wilson_interval(wins, n)
                    computed = (n / 5239, wins / n if n else math.nan, low / scale * 50)
                    record = {"kind": kind, "age_index": age_index, "card_name": name,
                              "is_absent_placeholder": name == EMPTY_CARD, "usage_count": n,
                              "wins": wins, "historical_score_scale": scale}
                    for index, field in enumerate(("selection_rate", "win_rate", "historical_score"), start=1):
                        cell = cached.cell(row_number, base + index)
                        if not isinstance(cell.value, (float, int)):
                            raise ValueError("Historical saved statistic is unavailable at {}".format(cell.coordinate))
                        value = computed[index - 1]
                        tolerance = 1e-8 if field == "historical_score" else 1e-12
                        record[field + "_python"] = value
                        record[field + "_excel"] = cell.value
                        record[field + "_difference"] = value - cell.value
                        record[field + "_matches"] = math.isclose(value, cell.value, rel_tol=0, abs_tol=tolerance)
                        record[field + "_cell"] = cell.coordinate
                    result.append(record)
    finally:
        formulas_book.close()
        cached_book.close()
    if len(result) != 84 or sum(row["is_absent_placeholder"] for row in result) != 4:
        raise ValueError("Expected 80 actual cards and four absent-leader rows in historical statistics")
    mismatched = sum(not row[field + "_matches"] for row in result
                     for field in ("selection_rate", "win_rate", "historical_score"))
    summary = {"source_file": "source_workbooks/5239.xlsx", "source_sheet": "統計",
               "openpyxl_version": openpyxl.__version__,
               "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
               "reference_games": 5239, "compared_rows": 84, "actual_card_rows": 80,
               "absent_leader_rows": 4, "compared_values": 252, "mismatched_values": mismatched,
               "historical_score_scales": sorted({row["historical_score_scale"] for row in result}),
               "note": "Compared preserved formula text and saved Excel values with Python recomputation; "
                       "the workbook was not edited or recalculated in native Excel."}
    return pd.DataFrame(result), summary
